import pandas as pd
from fastapi import UploadFile
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from core.config import get_connection_pool
from core.databases.gaji_batch_master_proses import fetch_gaji_batch_master_proses_by_root_batch_id
from core.proses_gaji import additional_gaji
# noinspection PyUnresolvedReferences
import swifter


def fetch_gaji_batch_master(batch_root_id: str):
    query = "SELECT * FROM gaji_batch_master WHERE batch_root_id = %s"
    with get_connection_pool() as conn:
        with conn.cursor() as cursor:
            cursor.execute(query, (batch_root_id,))
            result = cursor.fetchall()
            return pd.DataFrame(result)


def insert_gaji_batch_master_proses(data):
    query = """
            INSERT INTO gaji_batch_master_proses
            (jenis_gaji, formula, kode, nama, nilai, nilai_formula, urut, batch_master_id)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
            """
    with get_connection_pool() as conn:
        with conn.cursor() as cursor:
            cursor.executemany(query, data)
            conn.commit()


def process_potongan_from_excel(root_batch_id: str, workbook: Workbook, gbm: pd.DataFrame):
    query_data = []
    exclude_sheet = ["Sheet1"]
    for sheet_name in workbook.sheetnames:
        if sheet_name in exclude_sheet:
            continue
        worksheet = workbook[sheet_name]
        query_data.extend(process_worksheet(worksheet, gbm))

    insert_gaji_batch_master_proses(query_data)
    gbp = pd.DataFrame(fetch_gaji_batch_master_proses_by_root_batch_id(root_batch_id))
    gbm = gbm.swifter.apply(lambda x: additional_gaji.recalculate(x, gbp), axis=1)
    additional_gaji.update_additional(gbm)
    workbook.close()


def process_worksheet(worksheet: Worksheet, gbm: pd.DataFrame) -> list:
    """Process a worksheet and generate INSERT query data for `gaji_batch_master_proses` table."""
    # Get headers from the worksheet
    headers = [cell.value for cell in worksheet[10]]
    headers[0] = "no"
    headers[1] = "nama"
    headers[2] = "nipam"
    headers[3] = "gaji_pokok"
    headers[-2] = "total_potongan"
    headers[-1] = "gaji_bersih"

    # Get data from the worksheet
    data = []
    for row in worksheet.iter_rows(min_row=11, max_row=worksheet.max_row - 1, values_only=True):
        if row[0] is None:
            continue
        data.append(row)

    # Convert data to a Pandas DataFrame
    df = pd.DataFrame(data, columns=headers)
    df = df.astype({"nipam": str, "nama": str})
    float_columns = df.columns.difference(["nipam", "nama"])
    df[float_columns] = df[float_columns].astype(float, copy=False)

    # Process data and generate INSERT query data
    return processing_data(gbm, df)


def processing_data(gaji_batch_master: pd.DataFrame, dataframe: pd.DataFrame) -> list:
    """Generate INSERT query data for `gaji_batch_master_proses` table."""
    query_data = []
    for _, row in dataframe.iterrows():
        matching_gbm = gaji_batch_master.query(f"nipam == '{row['nipam']}'")
        if not matching_gbm.empty:
            row["batch_master_id"] = matching_gbm["id"].values[0]
            insert_data = generate_insert_query_data(row)
            query_data.extend(insert_data)
    return query_data


def generate_insert_query_data(row: pd.Series) -> list:
    insert_data = []
    excluded_columns = {"no", "nama", "nipam", "gaji_pokok",
                        "total_potongan", "gaji_bersih", "batch_master_id"}

    for column, value in row.items():
        if not pd.isna(value) and value != 0.0 and column not in excluded_columns:
            entry = (
                "POTONGAN",
                "",
                f"ADD_{column}",
                column,
                value,
                value,
                99,
                row["batch_master_id"]
            )
            insert_data.append(entry)

    return insert_data


async def read_excel(root_batch_id: str, file: UploadFile):
    content = await file.read()
    tmp_file_path = f"/tmp/{file.filename}"
    with open(tmp_file_path, "wb") as tmp_file:
        tmp_file.write(content)

    workbook = load_workbook(tmp_file_path, data_only=True)
    gaji_batch_master = fetch_gaji_batch_master(root_batch_id)
    gaji_batch_master = gaji_batch_master.astype({"nipam": str})

    process_potongan_from_excel(root_batch_id, workbook, gaji_batch_master)
