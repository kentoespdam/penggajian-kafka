from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from core.config import LOGGER
from core.helpers import get_previous_period, extract_period_year_month
from core.models.gaji_batch_master import fetch_daftar_potongan_gaji_by_batch_root_id
from core.models.gaji_batch_master_proses import fetch_additional_gaji_batch_master_proses_by_periode
from core.models.organisasi import fetch_organisasi_by_level
from core.process_gaji.phase4_generate import generate_potongan
from core.process_gaji.phase4_helper import phase4_filter_add_potongan_in_nipam, phase4_compute_potongan_metadata

# Introduce constants for configuration and paths
ORG_LEVEL = 4
PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
TEMPLATE_PATH = PROJECT_ROOT / "excel_template" / "potongan_gaji_template.xlsx"
RESULT_DIR = PROJECT_ROOT / "result_excel"


def build_potongan_gaji(batch_root_id: str) -> None:
    """
    Phase 4 entry point: build potongan gaji Excel per organisasi for the given batch root ID.
    """
    start_time = datetime.now()
    LOGGER.info(f"Starting phase4: build potongan gaji for batch ID {batch_root_id}")

    organisasi_df = fetch_organisasi_by_level(ORG_LEVEL)
    if organisasi_df.empty:
        LOGGER.error("Organisasi not found (level %s). Abort.", ORG_LEVEL)
        return

    potongan_df = fetch_daftar_potongan_gaji_by_batch_root_id(batch_root_id)
    if potongan_df.empty:
        LOGGER.error("Daftar potongan gaji not found for batch ID %s. Abort.", batch_root_id)
        return

    periode, tahun, bulan = extract_period_year_month(batch_root_id)
    prev_periode = get_previous_period(periode)
    add_potongan_prev_df = fetch_additional_gaji_batch_master_proses_by_periode(prev_periode)

    _generate_excel(
        batch_root_id=batch_root_id,
        tahun=tahun,
        bulan=bulan,
        organisasi_df=organisasi_df,
        daftar_potongan_gaji_df=potongan_df,
        add_potongan_prev_df=add_potongan_prev_df,
    )

    duration = datetime.now() - start_time
    LOGGER.info(f"Phase4 complete: potongan gaji built for batch ID {batch_root_id} in {duration}")


def _generate_excel(
        batch_root_id: str,
        tahun: int,
        bulan: int,
        organisasi_df: pd.DataFrame,
        daftar_potongan_gaji_df: pd.DataFrame,
        add_potongan_prev_df: pd.DataFrame,
) -> None:
    """
    Generate the potongan gaji Excel workbook:
      - one worksheet per organisasi
      - dynamic potongan columns computed from previous period additions
    """
    wb = load_workbook(TEMPLATE_PATH)

    # Compute metadata for dynamic potongan columns
    add_kode_nama_df, max_col = phase4_compute_potongan_metadata(add_potongan_prev_df)

    for _, organisasi in organisasi_df.iterrows():
        mask = daftar_potongan_gaji_df["kode_organisasi"].str.startswith(f"{organisasi['kode']}")
        potongan_gaji_pegawai_df = daftar_potongan_gaji_df[mask].reset_index(drop=True)
        add_potongan_df = phase4_filter_add_potongan_in_nipam(
            add_potongan_prev_df,
            potongan_gaji_pegawai_df["nipam"],
        )

        generate_potongan(
            wb=wb,
            title=f"{organisasi['nama']}",
            short_name=f"{organisasi['short_name']}",
            tahun=tahun,
            bulan=bulan,
            daftar_gaji_pegawai=potongan_gaji_pegawai_df,
            add_potongan_df=add_potongan_df,
            add_kode_nama_df=add_kode_nama_df,
            max_col=max_col,
        )

    # Remove template sheet and save
    if "Sheet1" in wb.sheetnames:
        wb.remove(wb["Sheet1"])

    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    output_path = RESULT_DIR / f"potongan_gaji_{batch_root_id}.xlsx"
    wb.save(output_path)
