import itertools

import pandas as pd
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from core.excel_helper import cell_builder, BORDER_TLR_THIN, BORDER_LRB_THIN, BORDER_THIN, BORDER_LR_THIN, auto_width
from core.process_gaji.phase4_helper import (phase4_prepare_sheet, PositionTracker, TITLE_LINES, FIRST_POTONGAN_COL,
                                             HEADER_ROW, SUBHEADER_ROW, phase4_get_nilai_by_kode)


def generate_potongan(
        wb: Workbook,
        title: str,
        short_name: str,
        tahun: int,
        bulan: int,
        daftar_gaji_pegawai: pd.DataFrame,
        add_potongan_df: pd.DataFrame,
        add_kode_nama_df: pd.DataFrame,
        max_col: int
):
    # Prepare worksheet and titles
    worksheet = phase4_prepare_sheet(wb, short_name, title, tahun, bulan)

    # Write title block and table headers
    _write_title_block(worksheet, max_col)
    _write_table_headers(worksheet, add_kode_nama_df)

    # Data rows
    tracker = PositionTracker(row=11, order=1)
    for _, pegawai_row in daftar_gaji_pegawai.iterrows():
        _generate_empty_row(worksheet, tracker.next_row(), max_col, border=BORDER_TLR_THIN)
        _generate_potongan_row(
            worksheet=worksheet,
            pos=tracker,
            gaji_pegawai=pegawai_row,
            add_kode_nama_df=add_kode_nama_df,
            add_potongan_df=add_potongan_df,
        )
        _generate_empty_row(worksheet, tracker.next_row(), max_col, border=BORDER_LRB_THIN)


def _write_title_block(worksheet: Worksheet, end_col: int):
    for row_num, content, has_bottom_border in TITLE_LINES:
        _write_merged_centered(worksheet=worksheet, row_num=row_num, content=content, end_col=end_col,
                               border_bottom=has_bottom_border)


def _write_merged_centered(
        worksheet: Worksheet,
        row_num: int,
        content: str,
        end_col: int,
        border_bottom: bool = False,
) -> None:
    """Write a centered, merged row with optional bottom borders."""
    start_col = 1
    cell_builder(worksheet, row_num, start_col, content, border={"bottom": "double"} if border_bottom else None,
                 horizontal_alignment="center")
    worksheet.merge_cells(
        start_row=row_num, start_column=start_col, end_row=row_num, end_column=end_col
    )


def _write_table_headers(worksheet: Worksheet, add_kode_nama_df: pd.DataFrame):
    # number of dynamic "potongan tambahan" columns
    potongan_count = 1 if add_kode_nama_df.empty else add_kode_nama_df["kode"].size
    column_counter = itertools.count(start=FIRST_POTONGAN_COL)

    # Potongan group header (row 9), spanning dynamic columns
    cell_builder(
        worksheet,
        HEADER_ROW,
        FIRST_POTONGAN_COL,
        "Potongan",
        border=BORDER_THIN,
        horizontal_alignment="center",
    )
    worksheet.merge_cells(
        start_row=HEADER_ROW,
        start_column=FIRST_POTONGAN_COL,
        end_row=HEADER_ROW,
        end_column=FIRST_POTONGAN_COL + potongan_count - 1 if potongan_count > 0 else FIRST_POTONGAN_COL,
    )

    def _write_subheader_cell(content: str) -> None:
        curr_cell = cell_builder(
            worksheet,
            SUBHEADER_ROW,
            next(column_counter),
            content,
            border=BORDER_THIN,
            horizontal_alignment="center",
        )
        auto_width(worksheet, curr_cell, content)

    # Subheaders for dynamic potongan columns (row 10)
    if add_kode_nama_df.empty:
        _write_subheader_cell("- POTONGAN TAMBAHAN -")
    else:
        for row in add_kode_nama_df.itertuples():
            _write_subheader_cell(row.nama)

    # Static columns: "JUMLAH POTONGAN" and "GAJI BERSIH"
    total_potongan_col = FIRST_POTONGAN_COL + potongan_count
    gaji_bersih_col = total_potongan_col + 1
    _write_static_header(worksheet, "JUMLAH POTONGAN", total_potongan_col, padding=4)
    _write_static_header(worksheet, "GAJI BERSIH", gaji_bersih_col, padding=2)


def _write_static_header(worksheet: Worksheet, title: str, col_num: int, padding: int) -> None:
    cell = cell_builder(worksheet, HEADER_ROW, col_num, title, border=BORDER_THIN, horizontal_alignment="center", )
    worksheet.column_dimensions[cell.column_letter].width = len(title) + padding
    worksheet.merge_cells(start_row=HEADER_ROW, start_column=col_num, end_row=SUBHEADER_ROW, end_column=col_num)


def _generate_empty_row(worksheet: Worksheet, row_num: int, max_col: int, border: dict | None = None) -> None:
    for index in range(max_col):
        cell_builder(worksheet, row_num, index + 1, "", border=border if border else BORDER_LR_THIN)


def _generate_potongan_row(
        worksheet: Worksheet,
        pos: PositionTracker,
        gaji_pegawai: pd.Series,
        add_kode_nama_df: pd.DataFrame,
        add_potongan_df: pd.DataFrame,
):
    col_num = itertools.count(start=1)
    current_row = pos.next_row()

    def build_cell(content: str | int | float, is_number: bool = False,
                   border: dict | None = None):
        cell = cell_builder(
            worksheet,
            current_row,
            next(col_num),
            content,
            border=border if border else BORDER_LR_THIN,
        )
        if is_number:
            cell.number_format = "#,##0"
        return cell

    # Identity and base salary
    build_cell(pos.next_order())
    build_cell(f"{gaji_pegawai.nama}")
    build_cell(f"{gaji_pegawai.nipam}")
    gaji_cell = build_cell(gaji_pegawai.gaji, is_number=True)

    total_potongan = 0
    first_cell = None
    last_cell = None
    if add_kode_nama_df.empty:
        first_cell = build_cell(0, is_number=True)
        last_cell = first_cell
    else:
        index_counter = itertools.count(start=0)
        for row in add_kode_nama_df.itertuples():
            amount = phase4_get_nilai_by_kode(add_potongan_df, gaji_pegawai.nipam, row.kode)
            total_potongan += amount
            current_cell = build_cell(amount, is_number=True)
            if next(index_counter) == 0:
                first_cell = current_cell
            last_cell = current_cell

    total_potongan_cell = build_cell(0, True)
    if first_cell and last_cell:
        total_potongan_cell.value = f"=SUM({get_column_letter(first_cell.column)}{first_cell.row}:{get_column_letter(last_cell.column)}{last_cell.row})"

    gaji_bersih_cell = build_cell(0, True)
    gaji_bersih_cell.value = f"={gaji_cell.coordinate}-{total_potongan_cell.coordinate}"
