from dataclasses import dataclass
from typing import Dict, Iterable
import pandas as pd
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from core.helpers import get_nama_bulan

# Introduce constant: centralized configuration for title lines
TITLE_LINES = [
    (1, "PEMERINTAH KABUPATEN BANYUMAS", False),
    (2, "PERUSAHAAN UMUM DAERAH AIR MINUM TIRTA SATRIA", False),
    (3, "Jalan Prof. DR. Suharso No. 52 Purwokerto 53114", False),
    (4, "Telp. & Fax. (0281)-632324", True),
    (6, "DAFTAR POTONGAN GAJI PEGAWAI", False),
]

# Header configuration constants
HEADER_ROW = 9
SUBHEADER_ROW = 10
FIRST_POTONGAN_COL = 5

# Centralized required column sets
REQ_POTONGAN_SUM_COLS_EITHER = {"kode", "nilai"}
REQ_NIPAM_KODE_NILAI_COLS = {"nipam", "kode", "nilai"}


@dataclass
class PositionTracker:
    row: int
    order: int

    def next_row(self) -> int:
        """Return current row and advance internal pointer by 1."""
        current = self.row
        self.row += 1
        return current

    def next_order(self) -> int:
        """Return the current order and advance internal pointers by 1."""
        current = self.order
        self.order += 1
        return current


def phase4_prepare_sheet(wb: Workbook, short_name: str, title: str, tahun: int, bulan: int) -> Worksheet:
    """
    Prepare a worksheet copy with title and period headers.
    - wb: source workbook with a pre-formatted active template sheet
    - short_name: sheet title (worksheet name)
    - title: section title to be displayed
    - tahun/bulan: period info to be displayed
    """
    ws = wb.copy_worksheet(wb.active)
    ws.title = short_name

    bulan_cell = ws.cell(row=7, column=1)
    bulan_cell.value = f"Bulan: {get_nama_bulan(bulan)} {tahun}"
    ws.merge_cells(start_row=7, start_column=1, end_row=7, end_column=2)

    title_cell = ws.cell(row=8, column=1)
    title_cell.value = title
    ws.merge_cells(start_row=8, start_column=1, end_row=8, end_column=8)

    return ws


def phase4_compute_potongan_metadata(add_potongan_df: pd.DataFrame) -> tuple[pd.DataFrame, int]:
    """
    Compute metadata for dynamic potongan columns.
    Returns:
      - DataFrame of unique (kode, nama) pairs
      - max_col: total columns to be prepared for the sheet
    """
    is_empty = add_potongan_df.empty
    if is_empty:
        # Preserve original behavior: when empty, return the same schema and fixed width
        add_kode_nama_df = pd.DataFrame(columns=["kode", "nama"])
        return add_kode_nama_df, 7

    add_kode_nama_df = add_potongan_df[["kode", "nama"]].drop_duplicates()
    # 3 fixed identity columns (No, Nama, NIPAM) + 1 gaji + dynamic potongan + 2 totals
    dynamic_count = add_kode_nama_df["kode"].size
    max_col = 4 + dynamic_count + 2
    return add_kode_nama_df, max_col


def phase4_filter_add_potongan_in_nipam(add_potongan_df: pd.DataFrame, nipam: Iterable[str]) -> pd.DataFrame:
    """
    Filter additional potongan rows to those matching provided nipam values.
    Returns a copy with reset index; empty with consistent columns if input is empty.
    """
    if add_potongan_df.empty:
        return pd.DataFrame(columns=["id", "batch_master_id", "kode", "nama", "nilai", "nipam"])
    mask = add_potongan_df["nipam"].isin(nipam)
    return add_potongan_df[mask].reset_index(drop=True).copy()


def phase4_get_nilai_by_kode(add_potongan_df: pd.DataFrame, nipam: str, kode: str) -> float:
    """
    Sum nilai for a given nipam and kode.
    - Returns 0.0 when DataFrame is empty or required columns are missing.
    """
    if add_potongan_df.empty:
        return 0.0

    if not REQ_NIPAM_KODE_NILAI_COLS.issubset(add_potongan_df.columns):
        return 0.0

    mask = add_potongan_df["nipam"].eq(nipam) & add_potongan_df["kode"].eq(kode)
    nilai_series = add_potongan_df.loc[mask, "nilai"]
    return nilai_series.sum() if not nilai_series.empty else 0.0
