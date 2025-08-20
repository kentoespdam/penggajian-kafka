# Python
import itertools
from datetime import datetime

import pandas as pd
from openpyxl.styles import Alignment
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from core.config import LOGGER
from core.helpers import get_nama_bulan
from core.process_gaji.phase3_generate_hg_filter import (
    filter_organisasi,
    filter_gaji_direksi,
    filter_gaji_pegawai,
    filter_gaji_kontrak,
    filter_komponen_by_batch_ids,
)
from core.process_gaji.phase3_generate_hg_helper import (
    generate_empty_row,
    generate_row,
    prepare_hg_worksheet, PositionTracker, KomponenData, COLUMN_COUNT, generate_section_title,
)

# Constants to improve readability
START_DATA_ROW = 7


def generate_hg_sheet(
        workbook: Workbook,
        organisasi_df: pd.DataFrame,
        year: int,
        month: int,
        gaji_pegawai_df: pd.DataFrame,
        komponen_gaji_df: pd.DataFrame,
) -> None:
    """
    Generate HG sheet (Himpunan Gaji).
    Prepares the worksheet and fills sections for GAPOK and Tunjangan using filtered data.
    """
    start_time = datetime.now()
    LOGGER.info(f"Starting phase3: Generate HG sheet for year {year} and month {month}")
    # Prepare worksheet and header
    worksheet = prepare_hg_worksheet(workbook, year, month)

    # Prepare data for HG sheet
    komponen = _prepare_komponen_data(organisasi_df, gaji_pegawai_df, komponen_gaji_df)

    # Track row and ordering in a single place
    tracker = PositionTracker(row=START_DATA_ROW, order=1)

    # GAPOK + Honor Kontrak section
    _generate_gapok_section(worksheet, komponen, tracker)

    # Tunjangan section
    _generate_tunjangan_section(worksheet, komponen, tracker)

    # Penghasilan Kotor Row
    _generate_penghasilan_kotor_section(worksheet, komponen, tracker)

    # Potongan Section
    _generate_potongan_section(worksheet, komponen, tracker)

    # Total Potongan Section
    _generate_jumlah_section(worksheet, komponen, tracker)

    # TTD Section
    _generate_ttd_section(worksheet, tracker, gaji_pegawai_df, year, month)

    elapsed_time = datetime.now() - start_time
    LOGGER.info(f"Finished phase3: Generate HG sheet for year {year} and month {month} in {elapsed_time}")


def _prepare_komponen_data(
        organisasi_df: pd.DataFrame,
        gaji_pegawai_df: pd.DataFrame,
        komponen_gaji_df: pd.DataFrame
) -> KomponenData:
    # Split organizations by CABANG vs non-CABANG
    organisasi_pusat_df = filter_organisasi(organisasi_df)
    organisasi_cabang_df = filter_organisasi(organisasi_df, True)

    # Filter employee groups
    gaji_direksi_df = filter_gaji_direksi(gaji_pegawai_df)
    gaji_pegawai_pusat_df = filter_gaji_pegawai(organisasi_pusat_df, gaji_pegawai_df)
    gaji_pegawai_cabang_df = filter_gaji_pegawai(organisasi_cabang_df, gaji_pegawai_df)
    gaji_kontrak_df = filter_gaji_kontrak(gaji_pegawai_df)
    gaji_kontrak_pusat_df = filter_gaji_pegawai(organisasi_pusat_df, gaji_kontrak_df)
    gaji_kontrak_cabang_df = filter_gaji_pegawai(organisasi_cabang_df, gaji_kontrak_df)

    # Filter components
    komponen_gaji_direksi_df = filter_komponen_by_batch_ids(komponen_gaji_df, gaji_direksi_df["id"])
    komponen_gaji_pegawai_pusat_df = filter_komponen_by_batch_ids(komponen_gaji_df, gaji_pegawai_pusat_df["id"])
    # Include Direksi into Pusat aggregation
    komponen_gaji_pegawai_pusat_df = pd.concat(
        [komponen_gaji_direksi_df, komponen_gaji_pegawai_pusat_df],
        ignore_index=True
    )
    komponen_gaji_pegawai_cabang_df = filter_komponen_by_batch_ids(komponen_gaji_df, gaji_pegawai_cabang_df["id"])
    komponen_gaji_kontrak_pusat_df = filter_komponen_by_batch_ids(komponen_gaji_df, gaji_kontrak_pusat_df["id"])
    komponen_gaji_kontrak_cabang_df = filter_komponen_by_batch_ids(komponen_gaji_df, gaji_kontrak_cabang_df["id"])

    return KomponenData(
        pegawai_pusat=komponen_gaji_pegawai_pusat_df,
        pegawai_cabang=komponen_gaji_pegawai_cabang_df,
        kontrak_pusat=komponen_gaji_kontrak_pusat_df,
        kontrak_cabang=komponen_gaji_kontrak_cabang_df,
    )


def _generate_gapok_section(worksheet: Worksheet, data: KomponenData, pos: PositionTracker) -> None:
    """
    Generate 'Gaji Pokok' and 'Honor Tenaga Kontrak' rows, followed by two spacer rows.
    """
    # Row Gaji Pokok
    generate_row(worksheet, pos.next_row(), data.pegawai_pusat, data.pegawai_cabang, "GP", "Gaji Pokok",
                 pos.next_order())

    # Row Honor Tenaga Kontrak
    generate_row(worksheet, pos.next_row(), data.kontrak_pusat, data.kontrak_cabang, "GP", "Honor Tenaga Kontrak",
                 pos.next_order())

    # Spacer rows
    for _ in range(2):
        generate_empty_row(worksheet, pos.next_row(), COLUMN_COUNT)


def _generate_tunjangan_section(worksheet: Worksheet, data: KomponenData, pos: PositionTracker) -> None:
    # Generate Tunjangan Title
    generate_section_title(worksheet, pos, "Tunjangan-tunjangan:")

    components = {
        "TUNJ_SI": "Istri/Suami",
        "TUNJ_ANAK": "Anak",
        "TUNJ_JABATAN": "Jabatan / Pelaksana",
        "TUNJ_BERAS": "Beras",
        "TUNJ_KK": "Kegiatan Kerja",
        "TUNJ_AIR": "Air",
        "TUNJ_PPH21": "P.Ph. Pasal 21",
        "PEMBULATAN": "Pembulatan"
    }
    # Generate Tunjangan Values
    for key, value in components.items():
        generate_row(worksheet, pos.next_row(), data.pegawai_pusat, data.pegawai_cabang, key, value, pos.next_order())

    generate_empty_row(worksheet, pos.next_row(), COLUMN_COUNT)


def _generate_penghasilan_kotor_section(worksheet: Worksheet, data: KomponenData, pos: PositionTracker) -> None:
    generate_row(worksheet, pos.next_row(), data.pegawai_pusat, data.pegawai_cabang, "PENGHASILAN_KOTOR",
                 "Jumlah Penghasilan Kotor", font_bold=True, with_pembulatan=True)
    generate_empty_row(worksheet, pos.next_row(), COLUMN_COUNT)


def _generate_potongan_section(worksheet: Worksheet, data: KomponenData, pos: PositionTracker) -> None:
    # Generate Potongan Title
    generate_section_title(worksheet, pos, "Potongan-potongan:")

    generate_row(worksheet, pos.next_row(), data.pegawai_pusat, data.pegawai_cabang, "POT_PENSIUN", "Iuran Pensiun",
                 pos.next_order())
    generate_row(worksheet, pos.next_row(), data.pegawai_pusat, data.pegawai_cabang, "POT_ASTEK", "Iuran Jamsostek",
                 pos.next_order())
    generate_row(worksheet, pos.next_row(), data.kontrak_pusat, data.kontrak_cabang, "POT_ASTEK",
                 "Iuran Jamsostek T. Kontrak", pos.next_order())
    generate_row(worksheet, pos.next_row(), data.pegawai_pusat, data.pegawai_cabang, "POT_RUDIN", "Sewa Rumah",
                 pos.next_order())
    generate_row(worksheet, pos.next_row(), data.pegawai_pusat, data.pegawai_cabang, "POT_JP", "Iuran Jamsostek (JPn)",
                 pos.next_order())
    generate_row(worksheet, pos.next_row(), data.kontrak_pusat, data.kontrak_cabang, "POT_JP",
                 "Iuran Jamsostek (JPn) T. Kontrak", pos.next_order())
    generate_row(worksheet, pos.next_row(), data.pegawai_pusat, data.pegawai_cabang, "POT_ASKES",
                 "Jmn Kes Nasional (JPn)", pos.next_order())
    generate_row(worksheet, pos.next_row(), data.kontrak_pusat, data.kontrak_cabang, "POT_ASKES",
                 "Jmn Kes Nasional (JPn) T. Kontrak", pos.next_order())

    generate_row(worksheet, pos.next_row(), data.pegawai_pusat, data.pegawai_cabang, "POT_TKK", "TKK", pos.next_order())
    generate_row(worksheet, pos.next_row(), data.kontrak_pusat, data.kontrak_cabang, "POT_TKK", "Honor",
                 pos.next_order())
    generate_row(worksheet, pos.next_row(), data.pegawai_pusat, data.pegawai_cabang, "POT_PPH21", "P.Ph. Pasal 21",
                 pos.next_order())
    generate_row(worksheet, pos.next_row(), data.pegawai_pusat, data.pegawai_cabang, "POT_PPH21_PEMBINA",
                 "P.Ph. Pasal 21 Bd Pembina dll", pos.next_order())


def _generate_jumlah_section(worksheet: Worksheet, data: KomponenData, pos: PositionTracker) -> None:
    generate_row(worksheet, pos.next_row(), data.pegawai_pusat, data.pegawai_cabang, "POTONGAN", "Jumlah Potongan",
                 font_bold=True)
    generate_row(worksheet, pos.next_row(), data.pegawai_pusat, data.pegawai_cabang, "PENGHASILAN_BERSIH",
                 "Jumlah Penghasilan Bersih", font_bold=True, with_pembulatan=True)


def _generate_ttd_section(worksheet: Worksheet, pos: PositionTracker, gaji_pegawai_df: pd.DataFrame, year: int,
                          month: int) -> None:
    row_counter = itertools.count(start=pos.row)
    next(row_counter)
    next(row_counter)

    # Tanggal
    tanggal_row = next(row_counter)
    tanggal_cell = worksheet.cell(row=tanggal_row, column=6)
    tanggal_cell.value = f"Purwokerto,      {get_nama_bulan(month)} {year}"
    tanggal_cell.alignment = Alignment(horizontal="center")
    worksheet.merge_cells(
        start_row=tanggal_row, start_column=6, end_row=tanggal_row, end_column=8
    )

    # Title Direksi
    title_direksi_row = next(row_counter)
    title_direksi_cell = worksheet.cell(row=title_direksi_row, column=6)
    title_direksi_cell.value = "DIREKSI PERUMDAM TIRTA SATRIA"
    title_direksi_cell.alignment = Alignment(horizontal="center")
    worksheet.merge_cells(
        start_row=title_direksi_row, start_column=6, end_row=title_direksi_row, end_column=8
    )

    # Title Mengetahui & Kabupaten
    title_mengetahui_row = next(row_counter)
    title_mengetahui_cell = worksheet.cell(row=title_mengetahui_row, column=2)
    title_mengetahui_cell.value = "Mengetahui,"
    title_mengetahui_cell.alignment = Alignment(horizontal="center")
    worksheet.merge_cells(
        start_row=title_mengetahui_row, start_column=2, end_row=title_mengetahui_row, end_column=4
    )
    title_kabupaten_cell = worksheet.cell(row=title_mengetahui_row, column=6)
    title_kabupaten_cell.value = "Kabupaten Banyumas"
    title_kabupaten_cell.alignment = Alignment(horizontal="center")
    worksheet.merge_cells(
        start_row=title_mengetahui_row, start_column=6, end_row=title_mengetahui_row, end_column=8
    )

    # Title Direktur
    title_direktur_row = next(row_counter)
    title_direktur_cell = worksheet.cell(row=title_direktur_row, column=2)
    title_direktur_cell.value = "Direktur Utama"
    title_direktur_cell.alignment = Alignment(horizontal="center")
    worksheet.merge_cells(
        start_row=title_direktur_row, start_column=2, end_row=title_direktur_row, end_column=4
    )
    title_direktur_cell = worksheet.cell(row=title_direktur_row, column=6)
    title_direktur_cell.value = "Direktur Umum"
    title_direktur_cell.alignment = Alignment(horizontal="center")
    worksheet.merge_cells(
        start_row=title_direktur_row, start_column=6, end_row=title_direktur_row, end_column=8
    )
    next(row_counter)
    next(row_counter)
    next(row_counter)

    mask_direktur_utama = gaji_pegawai_df["level_id"].eq(2)
    direktur_utama = gaji_pegawai_df[mask_direktur_utama].reset_index(drop=True)
    mask_direktur_umum = gaji_pegawai_df["level_id"].eq(4)
    direktur_umum = gaji_pegawai_df[mask_direktur_umum].reset_index(drop=True)

    # Nama Direktur
    nama_direktur_row = next(row_counter)
    nama_direktur_cell = worksheet.cell(row=nama_direktur_row, column=2)
    nama_direktur_cell.value = direktur_utama["nama"].values[0]
    nama_direktur_cell.alignment = Alignment(horizontal="center")
    worksheet.merge_cells(
        start_row=nama_direktur_row, start_column=2, end_row=nama_direktur_row, end_column=4
    )
    nama_direktur_cell = worksheet.cell(row=nama_direktur_row, column=6)
    nama_direktur_cell.value = direktur_umum["nama"].values[0]
    nama_direktur_cell.alignment = Alignment(horizontal="center")
    worksheet.merge_cells(
        start_row=nama_direktur_row, start_column=6, end_row=nama_direktur_row, end_column=8
    )

    # Nipam Direktur
    nipam_direktur_row = next(row_counter)
    nipam_direktur_cell = worksheet.cell(row=nipam_direktur_row, column=2)
    nipam_direktur_cell.value = "NIPAM. " + direktur_utama["nipam"].values[0]
    nipam_direktur_cell.alignment = Alignment(horizontal="center")
    worksheet.merge_cells(
        start_row=nipam_direktur_row, start_column=2, end_row=nipam_direktur_row, end_column=4
    )
    nipam_direktur_cell = worksheet.cell(row=nipam_direktur_row, column=6)
    nipam_direktur_cell.value = "NIPAM. " + direktur_umum["nipam"].values[0]
    nipam_direktur_cell.alignment = Alignment(horizontal="center")
    worksheet.merge_cells(
        start_row=nipam_direktur_row, start_column=6, end_row=nipam_direktur_row, end_column=8
    )
