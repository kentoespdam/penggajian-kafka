import itertools

import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from core.enums import StatusPegawai
from core.excel_helper import cell_builder
from core.helpers import get_nama_bulan


def generate_hgpkp_sheet(
        workbook: Workbook,
        organisasi_df: pd.DataFrame,
        year: int,
        month: int,
        gaji_pegawai_df: pd.DataFrame,
        komponen_gaji_df: pd.DataFrame,
):
    workbook.active = workbook["HGPKP1"]
    worksheet = workbook.copy_worksheet(workbook.active)
    worksheet.title = "HGPKP"
    worksheet.cell(row=7, column=2, value=f"Bulan: {get_nama_bulan(month)} {year}")

    row_num = itertools.count(start=16)
    urut = itertools.count(start=2)

    # Generate direksi row
    mask = gaji_pegawai_df["level_id"].isin([2, 3, 4])
    direksi_ids = gaji_pegawai_df[mask]["id"].tolist()

    mask = komponen_gaji_df["batch_master_id"].isin(direksi_ids)
    komponen_gaji_direksi = komponen_gaji_df[mask].reset_index(drop=True)

    next_row = _generate_row(
        worksheet,
        next(row_num),
        komponen_gaji_direksi,
        len(direksi_ids),
        "DIREKSI",
        next(urut),
    )
    row_num = itertools.count(next_row)

    # Generate organisasi row
    mask = ~organisasi_df["nama"].str.startswith("CABANG")
    organisasi_wt_cabang = organisasi_df[mask].reset_index(drop=True)
    for _, organisasi in organisasi_wt_cabang.iterrows():
        mask_kode_organisasi = gaji_pegawai_df["kode_organisasi"].str.startswith(str(organisasi["kode"]))
        mask_status_pegawai = gaji_pegawai_df["status_pegawai"].ne(StatusPegawai.KONTRAK.value)
        mask = mask_kode_organisasi & mask_status_pegawai
        pegawai_ids = gaji_pegawai_df[mask]["id"].tolist()

        mask = komponen_gaji_df["batch_master_id"].isin(pegawai_ids)
        komponen_gaji_organisasi = komponen_gaji_df[mask].reset_index(drop=True)
        next_row = _generate_row(worksheet, next(row_num), komponen_gaji_organisasi, len(pegawai_ids),
                                 str(organisasi["nama"]),
                                 next(urut))
        row_num = itertools.count(next_row)

    # Generate footer
    _generate_footer(worksheet, next_row, gaji_pegawai_df, komponen_gaji_df, organisasi_wt_cabang)


def _generate_row(
        worksheet: Worksheet,
        row_num: int,
        salary_components: pd.DataFrame,
        jml_pegawai: int,
        row_name: str,
        urut: int | None = None,
        border_bottom: bool = False,
):
    row_counter = itertools.count(start=row_num)
    column_index = itertools.count(start=1)

    def build_cell(value, is_number=False, h_align=None, v_align=None):
        cell = cell_builder(
            worksheet=worksheet,
            row_num=row_num,
            column_num=next(column_index),
            content=value,
            horizontal_alignment=h_align,
            vertical_alignment=v_align,
            border={
                "top": "thin",
                "left": "thin",
                "right": "thin",
            },
        )
        if is_number:
            cell.number_format = "#,##0"
        return cell

    if urut:
        build_cell(urut, True)
        build_cell(row_name)
    else:
        build_cell(row_name, v_align="center")
        next(column_index)
    build_cell(jml_pegawai, h_align="center")
    columns = [
        ["GP", "0", "TUNJ_JABATAN", "TUNJ_AIR", "POT_PENSIUN", "POT_ASKES", "PENGHASILAN_BERSIH_FINAL", ""],
        ["TUNJ_SI", "0", "TUNJ_BERAS", "TUNJ_PPH21", "POT_ASTEK", "POT_TKK", "", ""],
        ["TUNJ_ANAK", "", "TUNJ_KK", "PENGHASILAN_KOTOR", "SEWA_RUDIN", "POT_PPH21", "", ""],
        ["JUMLAH", "", "TUNJ_KESEHATAN", "PEMBULATAN", "POT_JP", "POTONGAN", "", ""],
    ]
    for index, component_list in enumerate(columns):
        current_row = next(row_counter)
        _generate_cell_list(
            worksheet,
            current_row,
            salary_components,
            component_list,
            is_first=index == 0,
            border_bottom=border_bottom and index == 3,
        )

    return next(row_counter)


def _generate_cell_list(
        worksheet: Worksheet,
        row_num: int,
        salary_components: pd.DataFrame,
        komponen_list: list,
        is_first: bool = False,
        border_bottom: bool = False,
):
    col_num = itertools.count(start=4 if is_first else 1)
    border_top = False

    def build_cell(value: str | float, is_number: bool = False):
        cell = cell_builder(
            worksheet=worksheet,
            row_num=row_num,
            column_num=next(col_num),
            content=value,
            border={
                "top": "thin" if border_top else None,
                "left": "thin",
                "right": "thin",
                "bottom": "thin" if border_bottom else None
            },
        )
        if is_number:
            cell.number_format = "#,##0"
        return cell

    if not is_first:
        for _ in range(3):
            build_cell("")

    for component in komponen_list:
        border_top = False
        if is_first or component == "POTONGAN":
            border_top = True
        if component == "0":
            build_cell(0, True)
        elif component == "":
            build_cell("")
        elif component == "JUMLAH":
            gaji_pokok = _calculate_nilai_by_kode(salary_components, "GP")
            tunj_si = _calculate_nilai_by_kode(salary_components, "TUNJ_SI")
            tunj_anak = _calculate_nilai_by_kode(salary_components, "TUNJ_ANAK")
            jumlah = gaji_pokok + tunj_si + tunj_anak
            build_cell(jumlah, True)
        else:
            nilai = _calculate_nilai_by_kode(salary_components, component)
            build_cell(nilai, True)


def _generate_footer(worksheet: Worksheet, row_num: int, gaji_pegawai_df: pd.DataFrame, komponen_gaji_df: pd.DataFrame,
                     organisasi_wt_cabang: pd.DataFrame):
    mask_status_pegawai = gaji_pegawai_df["status_pegawai"].ne(StatusPegawai.KONTRAK.value)
    organisasi_kode_list = tuple(organisasi_wt_cabang["kode"].unique().tolist())
    mask_kode_organisasi = gaji_pegawai_df["kode_organisasi"].str.startswith(organisasi_kode_list)
    mask_level_id = gaji_pegawai_df["level_id"].isin([2, 3, 4])
    all_pegawai = gaji_pegawai_df[mask_status_pegawai & (mask_kode_organisasi | mask_level_id)]

    all_pegawai_ids = all_pegawai["id"].tolist()
    mask = komponen_gaji_df["batch_master_id"].isin(all_pegawai_ids)
    komponen_gaji_pegawai = komponen_gaji_df[mask].reset_index(drop=True)

    _generate_row(worksheet, row_num, komponen_gaji_pegawai, len(all_pegawai_ids),
                  row_name="Total Sampai Halaman Ini",
                  border_bottom=True)

    worksheet.merge_cells(start_row=row_num, start_column=1, end_row=row_num + 3, end_column=2)
    jml_cell = worksheet.cell(row=row_num, column=1)
    jml_cell.alignment = Alignment(horizontal="center", vertical="center")
    jml_cell.font = Font(bold=True)


def _calculate_nilai_by_kode(salary_components: pd.DataFrame, kode: str):
    mask = salary_components["kode"].eq(kode)
    filtered_components = salary_components[mask]
    return filtered_components["nilai"].sum()
